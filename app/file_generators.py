from __future__ import annotations

import csv
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Optional


MIME_BY_EXTENSION = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pdf": "application/pdf",
}

SUPPORTED_OUTPUT_EXTENSIONS = {".txt", ".md", ".csv", ".docx", ".xlsx"}


@dataclass(slots=True)
class GeneratedFile:
    filename: str
    content_type: str
    path: Path
    warning: Optional[str] = None


def generate_response_file(
    output_dir: Path,
    source_filename: str,
    content: str,
    file_id: str,
    output_extension: str | None = None,
) -> GeneratedFile:
    suffix = Path(source_filename).suffix.lower()
    stem = Path(source_filename).stem

    if output_extension:
        suffix = output_extension.lower()
        if not suffix.startswith("."):
            suffix = f".{suffix}"

    warning: Optional[str] = None
    if suffix == ".pdf":
        suffix = ".txt"
        warning = "PDF output is not supported: fallback to .txt"

    if suffix not in SUPPORTED_OUTPUT_EXTENSIONS:
        suffix = ".txt"

    output_name = f"{stem}_response{suffix}"
    destination = output_dir / f"{file_id}{suffix}"

    if suffix == ".txt":
        destination.write_text(content, encoding="utf-8")
    elif suffix == ".md":
        destination.write_text(content, encoding="utf-8")
    elif suffix == ".csv":
        with destination.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            for line in content.splitlines():
                writer.writerow([line])
    elif suffix == ".docx":
        try:
            from docx import Document  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is not installed") from exc
        document = Document()
        for line in content.splitlines() or [""]:
            document.add_paragraph(line)
        document.save(str(destination))
    elif suffix == ".xlsx":
        try:
            from openpyxl import Workbook  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("openpyxl is not installed") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Response"
        lines = content.splitlines() or [content]
        for idx, line in enumerate(lines, start=1):
            sheet.cell(row=idx, column=1, value=line)
        workbook.save(str(destination))
    elif suffix == ".pdf":
        try:
            from reportlab.lib.pagesizes import A4  # type: ignore
            from reportlab.pdfgen import canvas  # type: ignore
        except ImportError:
            destination = output_dir / f"{file_id}.txt"
            output_name = f"{stem}_response.txt"
            suffix = ".txt"
            warning = "reportlab is not available: fallback to .txt"
            destination.write_text(content, encoding="utf-8")
        else:
            pdf = canvas.Canvas(str(destination), pagesize=A4)
            width, height = A4
            y = height - 40
            for raw_line in content.splitlines() or [""]:
                wrapped = textwrap.wrap(raw_line, width=110) or [""]
                for line in wrapped:
                    pdf.drawString(40, y, line)
                    y -= 14
                    if y < 40:
                        pdf.showPage()
                        y = height - 40
            pdf.save()
    else:  # pragma: no cover
        destination.write_text(content, encoding="utf-8")

    return GeneratedFile(
        filename=output_name,
        content_type=MIME_BY_EXTENSION.get(suffix, "text/plain"),
        path=destination,
        warning=warning,
    )
